from pathlib import Path
import hashlib
import re
from bs4 import BeautifulSoup

TEXT_EXTS={".txt",".md",".html",".htm",".docx"}
IMAGE_EXTS={".png",".jpg",".jpeg"}


def extract_text(path):
    ext=path.suffix.lower()
    if ext in {".txt",".md"}:
        return path.read_text(encoding="utf-8",errors="replace"),{"type":ext[1:],"pages":None,"warnings":[],"modality":"text"}
    if ext in {".html",".htm"}:
        raw=path.read_text(encoding="utf-8",errors="replace")
        soup=BeautifulSoup(raw,"html.parser")
        for tag in soup(["script","style","noscript"]): tag.decompose()
        return soup.get_text("\n",strip=True),{"type":"html","pages":None,"warnings":[],"modality":"document"}
    if ext==".docx":
        from docx import Document
        doc=Document(path)
        parts=[]
        for p in doc.paragraphs:
            if p.text.strip(): parts.append(p.text)
        for ti,table in enumerate(doc.tables,1):
            parts.append(f"\n[TABLE {ti}]")
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts),{"type":"docx","pages":None,"warnings":[],"modality":"document"}
    if ext==".pdf":
        import pymupdf
        doc=pymupdf.open(path)
        pages=[]; scanned=[]
        for i,p in enumerate(doc):
            t=p.get_text("text")
            if not t.strip(): scanned.append(i+1)
            pages.append(f"[PAGE {i+1}]\n{t}")
        warnings=[]
        if scanned: warnings.append(f"Pages with no extractable text: {scanned}. Sarvam Document AI OCR is recommended for these pages.")
        return "\n\n".join(pages),{"type":"pdf","pages":len(doc),"warnings":warnings,"scanned_pages":scanned,"modality":"document"}
    if ext in IMAGE_EXTS:
        return "",{"type":"image","pages":1,"warnings":["Image input requires Sarvam Document AI OCR."],"modality":"image","needs_ocr":True}
    raise ValueError(f"Unsupported extension: {ext}")


def basic_inventory(path,text):
    return {"file_name":path.name,"extension":path.suffix.lower(),"size_bytes":path.stat().st_size,"sha256":hashlib.sha256(path.read_bytes()).hexdigest(),"source_characters":len(text),"source_lines":text.count("\n")+1 if text else 0}


def split_raw(text,max_chars=9000,overlap=700):
    text=text.replace("\x00","").strip()
    if len(text)<=max_chars: return [text] if text else []
    paragraphs=re.split(r"\n\s*\n",text)
    chunks=[]; cur=""
    for p in paragraphs:
        p=p.strip()
        if not p: continue
        if len(cur)+len(p)+2<=max_chars: cur=(cur+"\n\n"+p).strip()
        else:
            if cur: chunks.append(cur)
            tail=cur[-overlap:] if overlap else ""
            cur=(tail+"\n\n"+p).strip()
    if cur: chunks.append(cur)
    return chunks
